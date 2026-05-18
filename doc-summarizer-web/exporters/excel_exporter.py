import io
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "בית המשפט",
    "שופט",
    "תיאור המקרה",
    "פסק הדין במלל",
    "סכום לפני מע\"מ (₪)",
    "סכום אחרי מע\"מ (₪)",
    "שם קובץ",
]

COL_WIDTHS = [28, 20, 45, 55, 22, 22, 30]

HEADER_FILL   = PatternFill("solid", fgColor="1C1C1C")
HEADER_FONT   = Font(name="Arial", bold=True, color="FFC700", size=11)
ALT_FILL      = PatternFill("solid", fgColor="F7F7F7")
PLAIN_FILL    = PatternFill("solid", fgColor="FFFFFF")
BORDER_SIDE   = Side(style="thin", color="DDDDDD")
CELL_BORDER   = Border(left=BORDER_SIDE, right=BORDER_SIDE,
                       top=BORDER_SIDE, bottom=BORDER_SIDE)
AMOUNT_FONT   = Font(name="Arial", size=10, bold=True, color="1A6B2F")
ERR_FONT      = Font(name="Arial", size=10, italic=True, color="C0392B")


def _fmt_amount(val) -> str:
    if val is None:
        return "לא צוין"
    if val == 0:
        return "₪ 0"
    return f"₪ {val:,.2f}"


def export_to_excel(results: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "פסקי דין"
    ws.sheet_view.rightToLeft = True

    # Header row
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.fill      = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center",
                                   wrap_text=True)
        cell.border    = CELL_BORDER
    ws.row_dimensions[1].height = 28

    # Data rows
    for row_idx, r in enumerate(results, 2):
        fill = ALT_FILL if row_idx % 2 == 0 else PLAIN_FILL
        err  = bool(r.get("error"))

        def wc(col, value, is_amount=False):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.fill      = fill
            cell.border    = CELL_BORDER
            cell.alignment = Alignment(horizontal="right", vertical="top",
                                       wrap_text=True)
            if err:
                cell.font = ERR_FONT
            elif is_amount:
                cell.font = AMOUNT_FONT
            else:
                cell.font = Font(name="Arial", size=10)

        if err:
            wc(1, r.get("filename", ""))
            wc(2, f"שגיאה: {r.get('error', '')}")
            for c in range(3, 8):
                wc(c, "")
        else:
            wc(1, r.get("court") or "לא צוין")
            wc(2, r.get("judge") or "לא צוין")
            wc(3, r.get("case_description") or "לא צוין")
            wc(4, r.get("verdict") or "לא צוין")
            wc(5, _fmt_amount(r.get("amount_before_vat")), is_amount=True)
            wc(6, _fmt_amount(r.get("amount_after_vat")),  is_amount=True)
            wc(7, r.get("filename", ""))

        ws.row_dimensions[row_idx].height = 80

    # Column widths
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # Freeze header row
    ws.freeze_panes = "A2"

    # Summary sheet
    ws2 = wb.create_sheet("סיכום סכומים")
    ws2.sheet_view.rightToLeft = True

    valid = [r for r in results if not r.get("error") and r.get("amount_after_vat") is not None]
    total_before = sum(r.get("amount_before_vat") or 0 for r in valid)
    total_after  = sum(r.get("amount_after_vat")  or 0 for r in valid)

    summary_rows = [
        ("מספר פסקי דין שנותחו", len(results)),
        ("פסקי דין תקינים",      len(valid)),
        ("סה\"כ לפני מע\"מ",     _fmt_amount(total_before)),
        ("סה\"כ אחרי מע\"מ",     _fmt_amount(total_after)),
        ("תאריך הפקה",           str(datetime.date.today())),
    ]

    for r_idx, (label, value) in enumerate(summary_rows, 1):
        lc = ws2.cell(row=r_idx, column=1, value=label)
        lc.font      = Font(name="Arial", bold=True, size=11)
        lc.alignment = Alignment(horizontal="right")
        vc = ws2.cell(row=r_idx, column=2, value=value)
        vc.font      = Font(name="Arial", size=11)
        vc.alignment = Alignment(horizontal="right")

    ws2.column_dimensions["A"].width = 30
    ws2.column_dimensions["B"].width = 25

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
