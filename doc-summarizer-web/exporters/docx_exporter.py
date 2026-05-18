from docx import Document
from docx.oxml import OxmlElement
import io
import datetime


def set_rtl(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    bidi = OxmlElement("w:bidi")
    pPr.insert(0, bidi)


def export_to_docx(filename: str, summary: str) -> bytes:
    doc = Document()
    title = doc.add_heading(filename, level=1)
    set_rtl(title)
    sub = doc.add_paragraph(f"סיכום — {datetime.date.today()}")
    set_rtl(sub)
    doc.add_paragraph("")
    for line in summary.split("\n"):
        if line.strip():
            p = doc.add_paragraph(line.strip())
            set_rtl(p)
    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
