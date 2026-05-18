import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter


HEADERS = [
    "בית המשפט בו התקיים הדיון",
    "השופט ששפט באותו דיון",
    "תאור המקרה",
    "פסק הדין במלל",
    "הסכום שיש לליברה לשלם לפני מעמ",
    "הסכום שיש לליברה לשלם אחרי מעמ",
]

COL_WIDTHS = [16.0, 15.29, 14.43, 16.86, 15.0, 15.0]
HEADER_ROW_HEIGHT = 54.75
DATA_ROW_HEIGHT   = 27.0

BORDER_SIDE = Side(style="thin", color="000000")
CELL_BORDER = Border(left=BORDER_SIDE, right=BORDER_SIDE,
                     top=BORDER_SIDE, bottom=BORDER_SIDE)

HEADER_FONT = Font(name="Arial", size=11)
DATA_FONT   = Font(name="Arial", size=10)


def _fmt_amount(val) -> str:
    if val is None:
        return "לא צוין"
    if val == 0:
        return "0"
    return f"{val:,.2f}"


def export_to_excel(results: list) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "sheet1"

    # Header row
    for col_idx, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font      = HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center",
                                   wrap_text=True)
        cell.border    = CELL_BORDER
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT

    # Data rows
    for row_idx, r in enumerate(results, 2):
        err = bool(r.get("error"))

        def wc(col, value):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.font      = DATA_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center",
                                       wrap_text=True)
            cell.border    = CELL_BORDER

        if err:
            wc(1, r.get("filename", ""))
            wc(2, f"שגיאה: {r.get('error', '')}")
            for c in range(3, 7):
                wc(c, "")
        else:
            wc(1, r.get("court") or "")
            wc(2, r.get("judge") or "")
            wc(3, r.get("case_description") or "")
            wc(4, r.get("verdict") or "")
            wc(5, _fmt_amount(r.get("amount_before_vat")))
            wc(6, _fmt_amount(r.get("amount_after_vat")))

        ws.row_dimensions[row_idx].height = DATA_ROW_HEIGHT

    # Column widths
    for col_idx, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
